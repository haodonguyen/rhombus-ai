"""File browsing and lightweight previews.

Runs inside web requests, so it never uses Spark: CSV previews read only the first bytes
of the object, and XLSX previews are limited to files small enough to download.
"""

import csv
import io
from dataclasses import dataclass
from datetime import datetime
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.files.connections import S3Connection
from apps.files.exceptions import PreviewUnavailable, UnsupportedFileType
from apps.files.storage import get_s3_client, translate_s3_errors
from processing.file_types import FileType

SAMPLE_ROW_LIMIT = 10
CSV_PREVIEW_BYTES = 256 * 1024
XLSX_PREVIEW_MAX_BYTES = 50 * 1024 * 1024


@dataclass(frozen=True)
class StoredFile:
    key: str
    size: int
    last_modified: datetime
    file_type: FileType


@dataclass(frozen=True)
class FilePage:
    files: list[StoredFile]
    next_cursor: str | None


@dataclass(frozen=True)
class FilePreview:
    key: str
    file_type: FileType
    columns: list[str]
    sample_rows: list[list[str | None]]


def get_file_type(key: str) -> FileType:
    file_type = FileType.from_key(key)
    if file_type is None:
        raise UnsupportedFileType()
    return file_type


def verify(connection: S3Connection) -> None:
    """Check that the credentials work and can read the bucket, before storing them."""
    client = get_s3_client(connection)
    with translate_s3_errors():
        client.list_objects_v2(Bucket=connection.bucket, MaxKeys=1)


def list_files(
    connection: S3Connection, prefix: str = "", cursor: str | None = None, page_size: int = 100
) -> FilePage:
    """List supported files in key order.

    The cursor is the key of the last file returned; S3's `StartAfter` resumes after it,
    so pages have exactly `page_size` files even when unsupported objects are skipped.
    """
    client = get_s3_client(connection)
    files: list[StoredFile] = []
    start_after = cursor or ""
    exhausted = False

    with translate_s3_errors():
        while len(files) < page_size and not exhausted:
            params = {"Bucket": connection.bucket, "Prefix": prefix, "MaxKeys": page_size}
            if start_after:
                params["StartAfter"] = start_after
            response = client.list_objects_v2(**params)
            contents = response.get("Contents", [])

            for index, obj in enumerate(contents):
                file_type = FileType.from_key(obj["Key"])
                if file_type is not None:
                    files.append(
                        StoredFile(obj["Key"], obj["Size"], obj["LastModified"], file_type)
                    )
                if len(files) == page_size:
                    exhausted = index == len(contents) - 1 and not response.get("IsTruncated")
                    break
            else:
                exhausted = not response.get("IsTruncated") or not contents
                if contents:
                    start_after = contents[-1]["Key"]

    next_cursor = files[-1].key if files and not exhausted else None
    return FilePage(files=files, next_cursor=next_cursor)


def get_file_preview(connection: S3Connection, key: str) -> FilePreview:
    file_type = get_file_type(key)
    client = get_s3_client(connection)

    with translate_s3_errors(key):
        size = client.head_object(Bucket=connection.bucket, Key=key)["ContentLength"]
        if size == 0:
            return FilePreview(key=key, file_type=file_type, columns=[], sample_rows=[])

        if file_type is FileType.CSV:
            response = client.get_object(
                Bucket=connection.bucket, Key=key, Range=f"bytes=0-{CSV_PREVIEW_BYTES - 1}"
            )
            content = response["Body"].read()
            columns, rows = parse_csv_preview(content, truncated=size > len(content))
        else:
            if size > XLSX_PREVIEW_MAX_BYTES:
                raise PreviewUnavailable(
                    "This Excel file is too large to preview; its columns are checked "
                    "when the job runs."
                )
            content = client.get_object(Bucket=connection.bucket, Key=key)["Body"].read()
            columns, rows = parse_xlsx_preview(content)

    return FilePreview(key=key, file_type=file_type, columns=columns, sample_rows=rows)


def parse_csv_preview(content: bytes, truncated: bool) -> tuple[list[str], list[list[str | None]]]:
    text = content.decode("utf-8-sig", errors="replace")
    if truncated:
        # The byte range may end mid-row; drop the incomplete final line.
        text = text[: text.rfind("\n") + 1]
    reader = csv.reader(io.StringIO(text))
    header = next(reader, [])
    rows: list[list[str | None]] = [
        list(row) for _, row in zip(range(SAMPLE_ROW_LIMIT), reader, strict=False)
    ]
    return header, rows


def parse_xlsx_preview(content: bytes) -> tuple[list[str], list[list[str | None]]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise PreviewUnavailable("The Excel file could not be read.") from exc

    try:
        # Spark's Excel reader uses the first sheet, so the preview does too.
        rows_iter = workbook.worksheets[0].iter_rows(max_row=SAMPLE_ROW_LIMIT + 1, values_only=True)
        header_cells = list(next(rows_iter, ()))
        while header_cells and header_cells[-1] is None:
            header_cells.pop()
        header = [_cell_text(cell) or "" for cell in header_cells]
        rows = [[_cell_text(cell) for cell in row[: len(header)]] for row in rows_iter]
    finally:
        workbook.close()
    return header, rows


def _cell_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)
